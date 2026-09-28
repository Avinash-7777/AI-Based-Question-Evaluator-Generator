from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def load_pdf(file_path: str):

    reader = PdfReader(file_path)

    text_parts = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        text = page.extract_text()

        if text:

            text_parts.append(
                f"[Page {page_number}]\n{text}"
            )

    return "\n\n".join(text_parts)


def load_docx(file_path: str):

    document = Document(file_path)

    text_parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(
                    " | ".join(row_text)
                )

    return "\n".join(text_parts)


def load_pptx(file_path: str):

    presentation = Presentation(file_path)

    text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        if slide_text:

            text_parts.append(
                f"[Slide {slide_number}]\n"
                + "\n".join(slide_text)
            )

    return "\n\n".join(text_parts)


def load_txt(file_path: str):

    return Path(
        file_path
    ).read_text(
        encoding="utf-8"
    )


def load_document(file_path: str):

    extension = (
        Path(file_path)
        .suffix
        .lower()
    )

    if extension == ".pdf":

        return load_pdf(file_path)

    elif extension == ".docx":

        return load_docx(file_path)

    elif extension == ".pptx":

        return load_pptx(file_path)

    elif extension == ".txt":

        return load_txt(file_path)

    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )